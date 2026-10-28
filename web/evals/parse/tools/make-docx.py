# Builds the DOCX documents of the parse corpus from their truth and meta JSON.
#   python web/evals/parse/tools/make-docx.py      (needs python-docx)
# docx-plain: headings + bulleted lists; docx-table: a two-column table layout.
import glob, json, os
from docx import Document
from docx.shared import Pt

corpus = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "corpus")
DEFAULT_TITLES = {"summary": "Profile", "experience": "Experience", "projects": "Projects", "volunteering": "Volunteering",
                  "education": "Education", "skills": "Skills", "certifications": "Certifications", "interests": "Interests"}


def add_sections(doc, t, titles, skip=()):
    def h(key):
        doc.add_heading(titles.get(key, DEFAULT_TITLES[key]), level=2)

    if t.get("summary") and "summary" not in skip:
        h("summary")
        doc.add_paragraph(t["summary"])
    if t.get("experience") and "experience" not in skip:
        h("experience")
        for e in t["experience"]:
            p = doc.add_paragraph()
            p.add_run(e["role"]).bold = True
            p.add_run(f" · {e['company']}, {e['location']}\t{e['dates']}")
            for b in e["bullets"]:
                doc.add_paragraph(b, style="List Bullet")
    for key in ("projects", "volunteering"):
        if t.get(key) and key not in skip:
            h(key)
            for x in t[key]:
                if key == "projects":
                    doc.add_paragraph(f"{x['name']} — {x['description']} {x['tech']}")
                else:
                    doc.add_paragraph(f"{x['role']}, {x['org']}  {x['dates']}")
                    for b in x["bullets"]:
                        doc.add_paragraph(b, style="List Bullet")
    if t.get("education") and "education" not in skip:
        h("education")
        for e in t["education"]:
            doc.add_paragraph(f"{e['title']}, {e['org']}, {e['location']}  {e['year']}")
    if t.get("skills") and "skills" not in skip:
        h("skills")
        for s in t["skills"]:
            p = doc.add_paragraph()
            p.add_run(f"{s['category']}: ").bold = True
            p.add_run(", ".join(s["items"]))
    if t.get("certifications") and "certifications" not in skip:
        h("certifications")
        for c in t["certifications"]:
            doc.add_paragraph(f"{c['title']} — {c['org']} ({c['year']})")
    if t.get("interests") and "interests" not in skip:
        h("interests")
        doc.add_paragraph(" · ".join(t["interests"]))


def contact_lines(t):
    c = t["contact"]
    return [c["email"], c["phone"], c["location"]] + [f"{l['label']}: {l['url']}" for l in c.get("links", [])]


for meta_path in sorted(glob.glob(os.path.join(corpus, "*.meta.json"))):
    meta = json.load(open(meta_path, encoding="utf-8"))
    if not meta["file"].endswith(".docx"):
        continue
    t = json.load(open(os.path.join(corpus, f"{meta['id']}.truth.json"), encoding="utf-8"))
    titles = meta.get("titles", {})
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(t["contact"]["name"], level=1)
    if t.get("targetTitle"):
        doc.add_paragraph(t["targetTitle"])
    if meta["layout"] == "docx-table":
        left, right = doc.add_table(rows=1, cols=2).rows[0].cells
        left.text = "\n".join(contact_lines(t))
        for e in t.get("education", []):
            left.add_paragraph(f"{e['title']}, {e['org']}, {e['location']}  {e['year']}")
        for s in t.get("skills", []):
            left.add_paragraph(f"{s['category']}: {', '.join(s['items'])}")
        right.text = t.get("summary", "")
        for e in t.get("experience", []):
            right.add_paragraph(f"{e['role']} · {e['company']}, {e['location']}  {e['dates']}")
            for b in e["bullets"]:
                right.add_paragraph("• " + b)
        add_sections(doc, t, titles, skip=("summary", "experience", "education", "skills"))
    else:
        doc.add_paragraph(" | ".join(contact_lines(t)))
        add_sections(doc, t, titles)
    doc.core_properties.author = "hirecute evals"
    doc.save(os.path.join(corpus, meta["file"]))
    print(meta["file"])
